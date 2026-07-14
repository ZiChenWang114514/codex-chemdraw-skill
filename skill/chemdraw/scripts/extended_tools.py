"""Stable high-level wrappers for audited cdxml-toolkit workflows."""

from __future__ import annotations

from contextlib import redirect_stderr, redirect_stdout
import copy
import io
import json
import os
from pathlib import Path, PurePosixPath
import posixpath
import struct
import tempfile
from typing import Any, Optional
import xml.etree.ElementTree as ET
import zipfile

import artifact_safety
from runtime_diagnostics import diagnose_runtime


_CLEANUP_APPROACHES = {
    "bbox_center", "arrow_driven", "proportional", "compact",
    "golden_ratio", "chemdraw_mimic",
}
_ALIGN_MODES = {"rdkit", "rxnmapper", "kabsch"}
_CFB_SIGNATURE = bytes.fromhex("D0CF11E0A1B11AE1")
_CHEMDRAW_CLSID = bytes.fromhex("216DBA412EA0CE118FD90020AFD1F20C")
_CHEMDRAW_CLSID_TEXT = "41BA6D21-A02E-11CE-8FD9-0020AFD1F20C"


def _contract(outputs: dict[str, Any], warnings=None, metadata=None) -> dict[str, Any]:
    payload = {
        "ok": True,
        "outputs": outputs,
        "warnings": list(warnings or []),
        "metadata": dict(metadata or {}),
    }
    artifacts = artifact_safety.artifact_records(
        artifact_safety.paths_from_value(outputs)
    )
    if artifacts:
        payload["metadata"]["artifacts"] = artifacts
    return payload


def _source(path: str, *, suffixes: tuple[str, ...] | None = None) -> Path:
    source = Path(path).expanduser().resolve()
    if not source.is_file():
        raise FileNotFoundError(f"Input file does not exist: {source}")
    if suffixes and source.suffix.lower() not in suffixes:
        raise ValueError(f"Input must use one of {suffixes}: {source}")
    if source.suffix.lower() == ".cdxml":
        _validate_cdxml(source)
    elif source.suffix.lower() in {".pptx", ".docx"}:
        _validate_office_package(source)
    return source


def _directory(path: str) -> Path:
    directory = Path(path).expanduser().resolve()
    if not directory.is_dir():
        raise FileNotFoundError(f"Input directory does not exist: {directory}")
    return directory


def _destination(
    source: Path,
    output_path: str | None,
    *,
    tag: str,
    suffix: str | None = None,
) -> Path:
    return artifact_safety.resolve_destination(
        source=source,
        output_path=output_path,
        tag=tag,
        suffix=suffix or source.suffix,
    )


def _local_name(tag: str) -> str:
    return tag.rsplit("}", 1)[-1]


def _validate_cdxml(path: Path) -> None:
    try:
        root = ET.parse(path).getroot()
    except (ET.ParseError, OSError) as exc:
        raise ValueError(f"Invalid CDXML XML: {path}: {exc}") from exc
    if _local_name(root.tag) != "CDXML":
        raise ValueError(f"Invalid CDXML root element in {path}: {root.tag}")


def _validate_png(path: Path) -> tuple[int, int]:
    try:
        from PIL import Image

        with Image.open(path) as image:
            if image.format != "PNG":
                raise ValueError(
                    f"Expected PNG image data, found {image.format or 'unknown format'}"
                )
            image.verify()
        with Image.open(path) as image:
            image.load()
            dimensions = image.size
    except Exception as exc:
        raise ValueError(f"PNG could not be decoded: {path}: {exc}") from exc
    if dimensions[0] <= 0 or dimensions[1] <= 0:
        raise ValueError(f"PNG dimensions must be positive: {path}: {dimensions}")
    return dimensions


def _validate_svg(path: Path) -> None:
    try:
        root = ET.parse(path).getroot()
    except (ET.ParseError, OSError) as exc:
        raise ValueError(f"Invalid SVG XML: {path}: {exc}") from exc
    if _local_name(root.tag).lower() != "svg":
        raise ValueError(f"Invalid SVG root element in {path}: {root.tag}")


def _validate_chemdraw_ole_bytes(data: bytes, label: str) -> None:
    if not isinstance(data, (bytes, bytearray, memoryview)) or not data:
        raise ValueError(f"ChemDraw OLE data is empty: {label}")
    payload = bytes(data)
    if len(payload) < 1024 or payload[:8] != _CFB_SIGNATURE:
        raise ValueError(f"Invalid OLE compound-file header: {label}")
    if struct.unpack_from("<H", payload, 28)[0] != 0xFFFE:
        raise ValueError(f"Invalid OLE byte order: {label}")
    sector_shift = struct.unpack_from("<H", payload, 30)[0]
    if sector_shift not in {9, 12}:
        raise ValueError(f"Invalid OLE sector size: {label}")
    sector_size = 1 << sector_shift
    if (len(payload) - 512) % sector_size:
        raise ValueError(f"Truncated OLE compound file: {label}")
    if struct.unpack_from("<I", payload, 44)[0] == 0:
        raise ValueError(f"OLE compound file has no FAT sectors: {label}")
    if _CHEMDRAW_CLSID not in payload or b"ChemDraw" not in payload:
        raise ValueError(f"OLE object is not identified as ChemDraw: {label}")
    if "CONTENTS".encode("utf-16le") not in payload:
        raise ValueError(f"ChemDraw OLE CONTENTS stream is missing: {label}")
    try:
        import olefile

        if not olefile.isOleFile(io.BytesIO(payload)):
            raise ValueError(f"OLE data is not a parseable compound file: {label}")
        compound = olefile.OleFileIO(io.BytesIO(payload))
        try:
            streams = {parts[-1] for parts in compound.listdir(streams=True, storages=False)}
            required = {"\x01CompObj", "\x01Ole", "\x02OlePres000", "CONTENTS"}
            missing = sorted(required - streams)
            if missing:
                raise ValueError(f"ChemDraw OLE streams are missing {missing}: {label}")
            if str(compound.root.clsid).upper() != _CHEMDRAW_CLSID_TEXT:
                raise ValueError(f"OLE root CLSID is not ChemDraw: {label}")
            contents = compound.openstream("CONTENTS").read()
            if len(contents) < 8 or not contents.startswith(b"VjCD"):
                raise ValueError(f"ChemDraw OLE CONTENTS stream is not valid CDX: {label}")
            if b"ChemDraw" not in compound.openstream("\x01CompObj").read():
                raise ValueError(f"ChemDraw OLE CompObj stream is invalid: {label}")
        finally:
            compound.close()
    except ValueError:
        raise
    except Exception as exc:
        raise ValueError(f"ChemDraw OLE compound file cannot be parsed: {label}: {exc}") from exc


def _build_chemdraw_ole(cdx_data: bytes) -> bytes:
    """Build a valid CFB object despite the upstream mini-stream limitation."""
    payload = bytes(cdx_data)
    if not payload.startswith(b"VjCD"):
        raise ValueError("ChemDraw CDX data has an invalid header")
    from cdxml_toolkit.office.ole_embedder import build_ole_compound_file

    # cdxml-toolkit 0.5.17 does not populate the CFB mini stream. Padding CDX
    # to the regular-stream cutoff preserves the CDX prefix and trailing data.
    padded = payload.ljust(4096, b"\0")
    ole_data = build_ole_compound_file(padded)
    _validate_chemdraw_ole_bytes(ole_data, "generated ChemDraw OLE")
    return ole_data


def _relationship_target(relationship_part: str, target: str) -> str:
    target = target.replace("\\", "/")
    if target.startswith("/"):
        normalized = posixpath.normpath(target.lstrip("/"))
    else:
        rel_path = PurePosixPath(relationship_part)
        source_directory = rel_path.parent.parent
        normalized = posixpath.normpath(str(source_directory / target))
    if normalized == ".." or normalized.startswith("../"):
        raise ValueError(f"OOXML relationship escapes the package: {target}")
    return normalized.lstrip("./")


def _validate_office_package(
    path: Path,
    *,
    expected_ole_objects: int | None = None,
    minimum_chemdraw_objects: int = 0,
) -> dict[str, int]:
    try:
        with zipfile.ZipFile(path, "r") as package:
            member_list = package.namelist()
            if len(member_list) != len(set(member_list)):
                raise ValueError(f"OOXML package contains duplicate parts: {path}")
            members = set(member_list)
            required = {"[Content_Types].xml", "_rels/.rels"}
            required.add(
                "ppt/presentation.xml"
                if path.suffix.lower() == ".pptx"
                else "word/document.xml"
            )
            missing = sorted(required - members)
            if missing:
                raise ValueError(
                    f"OOXML package is missing required parts {missing}: {path}"
                )
            corrupt = package.testzip()
            if corrupt:
                raise ValueError(f"OOXML package has a corrupt part {corrupt}: {path}")

            ole_payloads: list[tuple[str, bytes]] = []
            for member in member_list:
                if not (member.endswith(".xml") or member.endswith(".rels")):
                    continue
                try:
                    root = ET.fromstring(package.read(member))
                except ET.ParseError as exc:
                    raise ValueError(
                        f"OOXML part is not valid XML: {member}: {exc}"
                    ) from exc
                root_name = _local_name(root.tag)
                if member == "[Content_Types].xml" and root_name != "Types":
                    raise ValueError(
                        "OOXML content types part has the wrong root element: "
                        f"{root.tag}"
                    )
                if member.endswith(".rels") and root_name != "Relationships":
                    raise ValueError(
                        f"OOXML relationships part has the wrong root element: {member}"
                    )
                if member == "ppt/presentation.xml" and root_name != "presentation":
                    raise ValueError("OOXML presentation part has the wrong root element")
                if member == "word/document.xml" and root_name != "document":
                    raise ValueError("OOXML document part has the wrong root element")
                if not member.endswith(".rels"):
                    continue
                for relationship in root:
                    if not relationship.attrib.get("Type", "").endswith("/oleObject"):
                        continue
                    if relationship.attrib.get("TargetMode") == "External":
                        raise ValueError(
                            f"OOXML OLE relationship must be embedded: {member}"
                        )
                    target = relationship.attrib.get("Target")
                    if not target:
                        raise ValueError(f"OOXML OLE relationship has no target: {member}")
                    target_part = _relationship_target(member, target)
                    if target_part not in members:
                        raise ValueError(
                            f"OOXML OLE relationship target is missing: {target_part}"
                        )
                    ole_payloads.append((target_part, package.read(target_part)))
    except (zipfile.BadZipFile, OSError) as exc:
        raise ValueError(f"Invalid OOXML ZIP package: {path}: {exc}") from exc

    try:
        if path.suffix.lower() == ".pptx":
            from pptx import Presentation

            Presentation(str(path))
        else:
            from docx import Document

            Document(str(path))
    except Exception as exc:
        raise ValueError(f"OOXML package cannot be opened by its document parser: {path}: {exc}") from exc

    if expected_ole_objects is not None and len(ole_payloads) != expected_ole_objects:
        raise ValueError(
            f"Expected {expected_ole_objects} embedded OLE object(s), "
            f"found {len(ole_payloads)}: {path}"
        )

    valid_chemdraw = 0
    for label, payload in ole_payloads:
        try:
            _validate_chemdraw_ole_bytes(payload, f"{path}!/{label}")
        except ValueError:
            if expected_ole_objects is not None:
                raise
            continue
        valid_chemdraw += 1

    required_chemdraw = (
        expected_ole_objects
        if expected_ole_objects is not None
        else minimum_chemdraw_objects
    )
    if valid_chemdraw < required_chemdraw:
        raise ValueError(
            f"Expected at least {required_chemdraw} valid ChemDraw OLE object(s), "
            f"found {valid_chemdraw}: {path}"
        )
    return {"ole_objects": len(ole_payloads), "chemdraw_objects": valid_chemdraw}


def _assert_output(
    path: Path,
    *,
    expected_ole_objects: int | None = None,
    minimum_chemdraw_objects: int = 0,
) -> Any:
    if not path.is_file() or path.stat().st_size == 0:
        raise RuntimeError(f"Expected output was not created: {path}")
    try:
        suffix = path.suffix.lower()
        if suffix == ".cdxml":
            _validate_cdxml(path)
        elif suffix == ".png":
            width, height = _validate_png(path)
            return {"width": width, "height": height}
        elif suffix == ".svg":
            _validate_svg(path)
        elif suffix in {".pptx", ".docx"}:
            return _validate_office_package(
                path,
                expected_ole_objects=expected_ole_objects,
                minimum_chemdraw_objects=minimum_chemdraw_objects,
            )
    except ValueError as exc:
        raise RuntimeError(f"Invalid generated output {path}: {exc}") from exc
    return None


def _publish_without_overwrite(source: Path, destination: Path) -> None:
    artifact_safety.publish_file(source, destination)


def _commit_staged_outputs(staged: list[tuple[Path, Path]]) -> None:
    artifact_safety.publish_files(staged)


def _normalize_warnings(value: Any) -> list[str]:
    if value is None:
        return []
    if isinstance(value, str):
        return [value]
    if isinstance(value, (list, tuple)):
        return [str(item) for item in value]
    return [str(value)]


def _run_cleanup(source: str, destination: str, approach: str) -> Any:
    from cdxml_toolkit.layout.reaction_cleanup import run_cleanup

    return run_cleanup(source, destination, approach=approach, verbose=False)


def _render_cdxml(source: str, destination: str, dpi: int = 300) -> str:
    from cdxml_toolkit.chemdraw.cdxml_to_image import cdxml_to_image

    return cdxml_to_image(source, destination, png_dpi=dpi)


def _staged_cdxml_outputs(destination: Path, render_preview: bool, writer) -> dict[str, str]:
    preview = destination.with_name(f"{destination.stem}_preview.png")
    if render_preview and preview.exists():
        raise ValueError(f"Refusing to overwrite an existing preview: {preview}")
    with tempfile.TemporaryDirectory(prefix=".cdxml-", dir=destination.parent) as stage_dir:
        stage = Path(stage_dir)
        temporary_cdxml = stage / destination.name
        writer(temporary_cdxml)
        _assert_output(temporary_cdxml)
        staged = [(temporary_cdxml, destination)]
        outputs = {"cdxml": str(destination)}
        if render_preview:
            temporary_preview = stage / preview.name
            _render_cdxml(str(temporary_cdxml), str(temporary_preview), 300)
            _assert_output(temporary_preview)
            staged.append((temporary_preview, preview))
            outputs["preview_png"] = str(preview)
        _commit_staged_outputs(staged)
    return outputs


def _unlinked_sequential_steps(schemes: list[Any]) -> list[int]:
    unlinked = []
    for index, (current, following) in enumerate(zip(schemes, schemes[1:]), start=1):
        products = {value for value in current.get_product_smiles_set() if value}
        reactants = {value for value in following.get_reactant_smiles_set() if value}
        if not products or not reactants or products.isdisjoint(reactants):
            unlinked.append(index)
    return unlinked


def _auto_plan_safety_issues(schemes: list[Any], plan: Any, classify_pair) -> list[str]:
    chain = list(plan.sequential_chain)
    if len(chain) < 2:
        return []
    chain_set = set(chain)
    edges: set[tuple[int, int]] = set()
    groups = list(plan.parallel_groups)
    for left_group in range(len(groups)):
        for right_group in range(left_group + 1, len(groups)):
            for left_index in groups[left_group]:
                for right_index in groups[right_group]:
                    classification = classify_pair(
                        schemes[left_index], schemes[right_index]
                    )
                    if classification == "sequential_ab":
                        edges.add((left_group, right_group))
                    elif classification == "sequential_ba":
                        edges.add((right_group, left_group))

    chain_edges = {edge for edge in edges if edge[0] in chain_set and edge[1] in chain_set}
    expected = set(zip(chain, chain[1:]))
    missing = sorted(expected - chain_edges)
    unexpected = sorted(chain_edges - expected)
    incoming = {group: 0 for group in chain}
    outgoing = {group: 0 for group in chain}
    for source, destination in chain_edges:
        outgoing[source] += 1
        incoming[destination] += 1

    issues = []
    if missing:
        issues.append(
            "missing adjacent chemical links "
            + ", ".join(f"group {source}->{destination}" for source, destination in missing)
        )
    if unexpected or any(value > 1 for value in incoming.values()) or any(
        value > 1 for value in outgoing.values()
    ):
        detail = ", ".join(
            f"group {source}->{destination}" for source, destination in unexpected
        ) or "multiple incoming or outgoing links"
        issues.append(f"non-linear branching or convergent links ({detail})")
    return issues


def clean_scheme_layout(
    input_path: str,
    output_path: Optional[str] = None,
    approach: str = "chemdraw_mimic",
    render_preview: bool = True,
) -> dict[str, Any]:
    """Clean an existing CDXML reaction layout without changing the source file."""
    if approach not in _CLEANUP_APPROACHES:
        raise ValueError(f"Unsupported approach: {approach}")
    source = _source(input_path, suffixes=(".cdxml",))
    destination = _destination(source, output_path, tag="cleaned")
    outputs = _staged_cdxml_outputs(
        destination,
        render_preview,
        lambda temporary: _run_cleanup(str(source), str(temporary), approach),
    )
    return _contract(outputs, metadata={"approach": approach})


def merge_reaction_schemes(
    input_paths: list[str],
    output_path: Optional[str] = None,
    mode: str = "auto",
    equiv_mode: str = "default",
    reference_cdxml: Optional[str] = None,
    allow_adjacent: bool = True,
    render_preview: bool = True,
    force_sequential: bool = False,
) -> dict[str, Any]:
    """Merge parallel, sequential, or unrelated CDXML reaction schemes."""
    if len(input_paths) < 2:
        raise ValueError("merge_reaction_schemes requires at least two input files")
    if mode not in {"auto", "parallel", "sequential", "adjacent"}:
        raise ValueError(f"Unsupported mode: {mode}")
    if equiv_mode not in {"default", "no-equiv", "equiv-range"}:
        raise ValueError(f"Unsupported equiv_mode: {equiv_mode}")
    if mode in {"sequential", "adjacent"} and equiv_mode != "default":
        raise ValueError(f"equiv_mode is not used for {mode} merges")
    sources = [_source(path, suffixes=(".cdxml",)) for path in input_paths]
    destination = _destination(sources[0], output_path, tag="merged")
    reference = str(_source(reference_cdxml, suffixes=(".cdxml",))) if reference_cdxml else None
    from cdxml_toolkit.cdxml_utils import write_cdxml
    from cdxml_toolkit.layout.scheme_merger import (
        adjacent_place, auto_detect, classify_pair, execute_merge_plan, parallel_merge,
        parse_scheme, sequential_merge,
    )

    schemes = [parse_scheme(str(path)) for path in sources]
    plan_description = mode
    warnings = []
    if mode == "auto":
        plan = auto_detect(schemes)
        plan_description = plan.describe()
        plan_issues = _auto_plan_safety_issues(schemes, plan, classify_pair)
        if plan_issues and not force_sequential:
            raise ValueError(
                "The auto-detected sequential plan is unsafe: "
                + "; ".join(plan_issues)
                + ". Use force_sequential=true only after manual review."
            )
        if plan_issues:
            warnings.append(
                "Auto merge was forced despite an unsafe sequential plan: "
                + "; ".join(plan_issues)
            )
        tree = execute_merge_plan(
            schemes, plan, equiv_mode=equiv_mode, ref_cdxml=reference,
            allow_adjacent=allow_adjacent,
        )
    elif mode == "parallel":
        tree = parallel_merge(schemes, equiv_mode=equiv_mode, strict=True)
    elif mode == "sequential":
        unlinked = _unlinked_sequential_steps(schemes)
        if unlinked and not force_sequential:
            labels = ", ".join(f"{index}->{index + 1}" for index in unlinked)
            raise ValueError(
                "Sequential steps are not chemically linked by canonical product/reactant "
                f"SMILES at: {labels}. Use force_sequential=true only after manual review."
            )
        if unlinked:
            warnings.append(
                "Sequential merge was forced despite missing product/reactant links at "
                + ", ".join(f"{index}->{index + 1}" for index in unlinked)
            )
        tree = sequential_merge(schemes, ref_cdxml=reference)
    else:
        tree = adjacent_place([scheme.tree for scheme in schemes])
    outputs = _staged_cdxml_outputs(
        destination,
        render_preview,
        lambda temporary: write_cdxml(tree, str(temporary)),
    )
    return _contract(
        outputs,
        warnings=warnings,
        metadata={"mode": mode, "plan": plan_description},
    )


def polish_reaction_scheme(
    input_path: str,
    output_path: Optional[str] = None,
    merge_conditions: bool = True,
    approach: str = "chemdraw_mimic",
    align_mode: str = "rdkit",
    eln_csv: Optional[str] = None,
    reference_cdxml: Optional[str] = None,
    render_preview: bool = True,
) -> dict[str, Any]:
    """Run the audited deterministic polishing pipeline on a CDXML scheme."""
    if approach not in _CLEANUP_APPROACHES:
        raise ValueError(f"Unsupported approach: {approach}")
    if align_mode not in _ALIGN_MODES:
        raise ValueError(f"Unsupported align_mode: {align_mode}")
    source = _source(input_path, suffixes=(".cdxml",))
    destination = _destination(source, output_path, tag="polished")
    csv_path = str(_source(eln_csv, suffixes=(".csv",))) if eln_csv else None
    reference = str(_source(reference_cdxml, suffixes=(".cdxml",))) if reference_cdxml else None
    from cdxml_toolkit.deterministic_pipeline.legacy.scheme_polisher_v2 import run_pipeline

    def write_polished(temporary: Path) -> None:
        run_pipeline(
            str(source), str(temporary), merge_conditions=merge_conditions,
            approach=approach, align_mode=align_mode, eln_csv=csv_path,
            ref_cdxml=reference, verbose=False,
        )

    outputs = _staged_cdxml_outputs(destination, render_preview, write_polished)
    return _contract(outputs, metadata={"approach": approach, "align_mode": align_mode})


def render_cdxml_files(
    input_paths: list[str],
    output_dir: Optional[str] = None,
    format: str = "png",
    dpi: int = 300,
) -> dict[str, Any]:
    """Render one or more CDXML files through native ChemDraw COM."""
    format = format.lower()
    if format not in {"png", "svg"}:
        raise ValueError("format must be 'png' or 'svg'")
    if dpi < 36 or dpi > 2400:
        raise ValueError("dpi must be between 36 and 2400")
    if not input_paths:
        raise ValueError("input_paths must not be empty")
    sources = [_source(path, suffixes=(".cdxml",)) for path in input_paths]
    directory = Path(output_dir).expanduser().resolve() if output_dir else sources[0].parent
    if directory.exists() and not directory.is_dir():
        raise ValueError(f"output_dir is not a directory: {directory}")

    destinations = [directory / f"{source.stem}.{format}" for source in sources]
    destination_keys = [os.path.normcase(str(path)) for path in destinations]
    if len(destination_keys) != len(set(destination_keys)):
        raise ValueError("Render inputs produce duplicate destination paths")
    conflicts = [path for path in destinations if path.exists()]
    if conflicts:
        raise ValueError(
            "Refusing to overwrite existing render destination(s): "
            + ", ".join(str(path) for path in conflicts)
        )

    directory_existed = directory.exists()
    directory.mkdir(parents=True, exist_ok=True)
    dimensions = {}
    try:
        with tempfile.TemporaryDirectory(prefix=".render-", dir=directory) as stage_dir:
            stage = Path(stage_dir)
            staged = []
            for source, destination in zip(sources, destinations):
                temporary = stage / destination.name
                _render_cdxml(str(source), str(temporary), dpi)
                validation = _assert_output(temporary)
                if format == "png":
                    dimensions[str(destination)] = validation
                staged.append((temporary, destination))
            _commit_staged_outputs(staged)
    except Exception:
        if not directory_existed:
            try:
                directory.rmdir()
            except OSError:
                pass
        raise

    metadata = {"format": format, "dpi": dpi}
    if dimensions:
        metadata["dimensions"] = dimensions
    return _contract(
        {"rendered": [str(path) for path in destinations]}, metadata=metadata
    )


def _nested_cdxml_slots(value: Any):
    if isinstance(value, dict):
        if str(value.get("type", "")).lower() == "cdxml":
            yield value
        for child in value.values():
            yield from _nested_cdxml_slots(child)
    elif isinstance(value, list):
        for child in value:
            yield from _nested_cdxml_slots(child)


def _prevalidate_office_manifest(manifest: Path) -> list[Path]:
    try:
        data = json.loads(manifest.read_text(encoding="utf-8"))
    except (json.JSONDecodeError, OSError) as exc:
        raise ValueError(f"Invalid Office manifest JSON: {manifest}: {exc}") from exc
    if not isinstance(data, dict):
        raise ValueError("Office manifest root must be a JSON object")

    base = manifest.parent.resolve()
    resolved = []
    for slot in _nested_cdxml_slots(data):
        raw_path = slot.get("file")
        if not isinstance(raw_path, str) or not raw_path.strip():
            raise ValueError("Every CDXML manifest slot requires a non-empty file path")
        relative = Path(raw_path)
        if relative.is_absolute() or relative.drive:
            raise ValueError(f"CDXML manifest paths must not be absolute: {raw_path}")
        if ".." in relative.parts:
            raise ValueError(f"CDXML manifest path traversal is not allowed: {raw_path}")
        candidate = (base / relative).resolve()
        try:
            candidate.relative_to(base)
        except ValueError as exc:
            raise ValueError(
                f"CDXML manifest path must remain inside the manifest directory: {raw_path}"
            ) from exc
        resolved.append(_source(str(candidate), suffixes=(".cdxml",)))
    return resolved


def fill_office_template(
    template_path: str,
    manifest_path: str,
    output_path: Optional[str] = None,
) -> dict[str, Any]:
    """Fill PPTX/DOCX text and editable ChemDraw placeholders from a manifest."""
    template = _source(template_path, suffixes=(".pptx", ".docx"))
    template_package = _validate_office_package(template)
    manifest = _source(manifest_path, suffixes=(".json",))
    requested_cdxml = _prevalidate_office_manifest(manifest)
    destination = _destination(template, output_path, tag="filled")
    from cdxml_toolkit.office.doc_from_template import main

    stdout, stderr = io.StringIO(), io.StringIO()
    with tempfile.TemporaryDirectory(prefix=".fill-", dir=destination.parent) as stage_dir:
        temporary = Path(stage_dir) / destination.name
        try:
            with redirect_stdout(stdout), redirect_stderr(stderr):
                exit_code = main([
                    "--template", str(template), "--manifest", str(manifest),
                    "--output", str(temporary), "--json",
                ])
        except SystemExit as exc:
            raise RuntimeError(
                f"Template filling exited unexpectedly with code {exc.code}"
            ) from exc
        if exit_code:
            raise RuntimeError(stderr.getvalue().strip() or "Template filling failed")
        try:
            details = json.loads(stdout.getvalue())
        except json.JSONDecodeError as exc:
            raise RuntimeError(f"Template filling returned invalid JSON: {exc}") from exc
        if not isinstance(details, dict):
            raise RuntimeError("Template filling JSON result must be an object")
        embedded = details.get("ole_objects")
        if isinstance(embedded, bool) or not isinstance(embedded, int):
            raise RuntimeError("Template filling did not report a valid OLE object count")
        if embedded != len(requested_cdxml):
            raise RuntimeError(
                f"Requested {len(requested_cdxml)} CDXML embed(s), "
                f"but the toolkit reported {embedded} embedded OLE object(s)"
            )
        package = _assert_output(
            temporary,
            minimum_chemdraw_objects=(
                template_package["chemdraw_objects"] + len(requested_cdxml)
            ),
        )
        expected_package_ole = template_package["ole_objects"] + len(requested_cdxml)
        if package["ole_objects"] != expected_package_ole:
            raise RuntimeError(
                f"Requested {len(requested_cdxml)} new OLE embed(s), but the "
                f"output package contains {package['ole_objects']} total OLE object(s); "
                f"expected {expected_package_ole} including the template baseline"
            )
        _commit_staged_outputs([(temporary, destination)])
    return _contract(
        {"office": str(destination)}, warnings=details.get("warnings", []),
        metadata={"slots_filled": details.get("slots_filled", 0), "ole_objects": details.get("ole_objects", 0)},
    )


def batch_embed_cdxml_in_office(
    cdxml_paths: list[str],
    output_path: str,
    margin_pt: float = 0,
) -> dict[str, Any]:
    """Create PPTX or DOCX containing editable ChemDraw OLE objects."""
    if not cdxml_paths:
        raise ValueError("cdxml_paths must not be empty")
    if margin_pt < 0:
        raise ValueError("margin_pt must be non-negative")
    sources = [_source(path, suffixes=(".cdxml",)) for path in cdxml_paths]
    destination = Path(output_path).expanduser().resolve()
    if destination.suffix.lower() not in {".pptx", ".docx"}:
        raise ValueError("output_path must end in .pptx or .docx")
    if destination.exists() or destination in sources:
        raise ValueError(f"Refusing to overwrite an existing file: {destination}")
    destination.parent.mkdir(parents=True, exist_ok=True)
    from cdxml_toolkit.office.ole_embedder import (
        batch_convert, build_docx, build_pptx,
        get_cdxml_content_size,
    )

    converted = batch_convert([str(path) for path in sources])
    if len(converted) != len(sources):
        raise RuntimeError("ChemDraw did not convert every CDXML input")
    items = []
    for index, item in enumerate(converted, start=1):
        if not isinstance(item, dict):
            raise RuntimeError(f"ChemDraw conversion {index} returned invalid data")
        cdx_data = item.get("cdx_data")
        emf_data = item.get("emf_data")
        if not isinstance(cdx_data, (bytes, bytearray, memoryview)) or not cdx_data:
            raise RuntimeError(f"ChemDraw conversion {index} returned empty CDX data")
        if not isinstance(emf_data, (bytes, bytearray, memoryview)) or not emf_data:
            raise RuntimeError(f"ChemDraw conversion {index} returned empty EMF data")
        width, height = get_cdxml_content_size(item["path"], margin_pt=margin_pt)
        if width <= 0 or height <= 0:
            raise RuntimeError(
                f"ChemDraw conversion {index} returned invalid content dimensions"
            )
        try:
            ole_data = _build_chemdraw_ole(bytes(cdx_data))
        except ValueError as exc:
            raise RuntimeError(str(exc)) from exc
        items.append({
            "ole_data": ole_data,
            "emf_data": bytes(emf_data), "width_emu": width,
            "height_emu": height, "name": item["name"],
        })
    with tempfile.TemporaryDirectory(prefix=".embed-", dir=destination.parent) as stage_dir:
        temporary = Path(stage_dir) / destination.name
        if destination.suffix.lower() == ".pptx":
            build_pptx(items, str(temporary))
        else:
            build_docx(items, str(temporary))
        _assert_output(temporary, expected_ole_objects=len(items))
        _commit_staged_outputs([(temporary, destination)])
    return _contract({"office": str(destination)}, metadata={"objects_embedded": len(items)})


def inspect_chemdraw_objects_in_office(
    input_path: str,
    output_dir: Optional[str] = None,
    render_previews: bool = True,
) -> dict[str, Any]:
    """Inventory editable ChemDraw objects and extract numbered CDXML previews."""
    source = _source(input_path, suffixes=(".pptx", ".docx"))
    destination = artifact_safety.resolve_directory_destination(
        source=source,
        output_dir=output_dir,
        tag="chemdraw_objects",
    )
    import office_objects

    with artifact_safety.staging_directory(destination) as staged:
        manifest = office_objects.write_inspection(
            source,
            staged,
            render_previews=render_previews,
        )
        manifest_path = staged / "manifest.json"
        artifact_safety.validate_artifact(manifest_path)
        for item in manifest["objects"]:
            _assert_output(staged / item["cdxml"])
            if item.get("preview_png"):
                _assert_output(staged / item["preview_png"])
        artifact_safety.publish_directory(staged, destination)

    cdxml_files = [str(destination / item["cdxml"]) for item in manifest["objects"]]
    previews = [
        str(destination / item["preview_png"])
        for item in manifest["objects"]
        if item.get("preview_png")
    ]
    return _contract(
        {
            "output_dir": str(destination),
            "manifest": str(destination / "manifest.json"),
            "cdxml_files": cdxml_files,
            "previews": previews,
        },
        metadata={
            "objects": len(manifest["objects"]),
            "source_sha256": manifest["source_sha256"],
        },
    )


def _replacement_destination(
    source: Path,
    output_path: str | None,
    *,
    render_pdf_preview: bool,
) -> tuple[Path, Path | None]:
    destination = _destination(source, output_path, tag="replaced")
    if not render_pdf_preview:
        return destination, None
    pdf = destination.with_suffix(".pdf")
    if output_path is not None and pdf.exists():
        raise ValueError(f"Refusing to overwrite an existing file: {pdf}")
    if output_path is None:
        base = destination
        index = 2
        while destination.exists() or pdf.exists():
            destination = base.with_name(f"{base.stem}_{index}{base.suffix}")
            pdf = destination.with_suffix(".pdf")
            index += 1
    return destination, pdf


def replace_chemdraw_objects_in_office(
    input_path: str,
    replacements_manifest: str,
    output_path: Optional[str] = None,
    render_pdf_preview: bool = True,
) -> dict[str, Any]:
    """Replace selected ChemDraw OLE contents and previews without moving them."""
    source = _source(input_path, suffixes=(".pptx", ".docx"))
    manifest_path = _source(replacements_manifest, suffixes=(".json",))
    source_package = _validate_office_package(source)
    import office_objects

    source_sha256, objects = office_objects.scan_office_objects(source)
    replacements = office_objects.load_replacement_manifest(
        manifest_path,
        source_sha256=source_sha256,
        objects=objects,
    )
    destination, pdf_destination = _replacement_destination(
        source,
        output_path,
        render_pdf_preview=render_pdf_preview,
    )

    from cdxml_toolkit.office.ole_embedder import batch_convert

    with office_objects.com_apartment():
        converted = batch_convert(
            [str(item["replacement_cdxml"]) for item in replacements]
        )
    if len(converted) != len(replacements):
        raise RuntimeError("ChemDraw did not convert every replacement CDXML")
    replacement_parts: dict[str, dict[str, bytes]] = {}
    for index, (replacement, converted_item) in enumerate(
        zip(replacements, converted), start=1
    ):
        if not isinstance(converted_item, dict):
            raise RuntimeError(f"ChemDraw replacement conversion {index} was invalid")
        cdx_data = converted_item.get("cdx_data")
        emf_data = converted_item.get("emf_data")
        if not isinstance(cdx_data, (bytes, bytearray, memoryview)) or not cdx_data:
            raise RuntimeError(
                f"ChemDraw replacement conversion {index} returned empty CDX data"
            )
        if not isinstance(emf_data, (bytes, bytearray, memoryview)) or not emf_data:
            raise RuntimeError(
                f"ChemDraw replacement conversion {index} returned empty EMF data"
            )
        try:
            ole_data = _build_chemdraw_ole(bytes(cdx_data))
        except ValueError as exc:
            raise RuntimeError(str(exc)) from exc
        record = replacement["record"]
        replacement_parts[replacement["object_id"]] = {
            "embedding_part": record["embedding_part"],
            "preview_part": record["preview_part"],
            "ole_data": ole_data,
            "emf_data": bytes(emf_data),
        }

    staged_outputs: list[tuple[Path, Path]] = []
    with tempfile.TemporaryDirectory(
        prefix=".replace-office-", dir=destination.parent
    ) as stage_dir:
        temporary_office = Path(stage_dir) / destination.name
        office_objects.rewrite_office_package(
            source,
            temporary_office,
            replacement_parts,
        )
        output_package = _assert_output(
            temporary_office,
            minimum_chemdraw_objects=source_package["chemdraw_objects"],
        )
        if output_package != source_package:
            raise RuntimeError(
                "Office replacement changed the package OLE object inventory: "
                f"before={source_package}, after={output_package}"
            )
        staged_outputs.append((temporary_office, destination))
        if pdf_destination is not None:
            temporary_pdf = Path(stage_dir) / pdf_destination.name
            office_objects.render_office_pdf(temporary_office, temporary_pdf)
            office_objects.validate_pdf(temporary_pdf)
            staged_outputs.append((temporary_pdf, pdf_destination))
        _commit_staged_outputs(staged_outputs)

    outputs = {"office": str(destination)}
    if pdf_destination is not None:
        outputs["pdf_preview"] = str(pdf_destination)
    return _contract(
        outputs,
        metadata={
            "objects_replaced": len(replacements),
            "object_ids": [item["object_id"] for item in replacements],
            "source_sha256": source_sha256,
        },
    )


def _discover_experiment(input_dir: str, experiment: str | None):
    from cdxml_toolkit.analysis.deterministic.discover_experiment_files import discover_experiment_files

    return discover_experiment_files(input_dir, experiment)


def discover_experiment_files(
    input_dir: str,
    experiment: Optional[str] = None,
    output_path: Optional[str] = None,
) -> dict[str, Any]:
    """Discover and classify files belonging to one experiment."""
    directory = _directory(input_dir)
    try:
        result = _discover_experiment(str(directory), experiment).to_dict()
    except SystemExit as exc:
        raise RuntimeError(
            f"Experiment discovery exited unexpectedly with code {exc.code}"
        ) from exc
    if not isinstance(result, dict):
        raise RuntimeError("Experiment discovery returned invalid data")
    payload = _contract(
        {"discovery": result},
        warnings=_normalize_warnings(result.get("warnings")),
        metadata={"experiment": result.get("experiment")},
    )
    if output_path:
        destination = _destination(directory / "experiment", output_path, tag="files", suffix=".json")
        payload["outputs"]["json"] = str(destination)
        with tempfile.TemporaryDirectory(prefix=".discovery-", dir=destination.parent) as stage_dir:
            temporary = Path(stage_dir) / destination.name
            temporary.write_text(json.dumps(payload, indent=2), encoding="utf-8")
            _assert_output(temporary)
            _commit_staged_outputs([(temporary, destination)])
        payload = artifact_safety.with_artifacts(payload, [destination])
    return payload


def analyze_lcms_series(
    files: list[str],
    output_path: Optional[str] = None,
    rt_tolerance: float = 0.02,
    mz_tolerance: float = 0.5,
    trend_threshold: float = 0.2,
    ignore_instrument: bool = False,
) -> dict[str, Any]:
    """Analyze a chronological series of standard LCMS PDF reports."""
    if len(files) < 2:
        raise ValueError("analyze_lcms_series requires at least two files")
    if min(rt_tolerance, mz_tolerance, trend_threshold) <= 0:
        raise ValueError("LCMS tolerances and trend_threshold must be positive")
    sources = [_source(path, suffixes=(".pdf",)) for path in files]
    anchor = sources[0]
    destination = _destination(anchor, output_path, tag="lcms-series", suffix=".json")
    from cdxml_toolkit.analysis.deterministic.multi_lcms_analyzer import main

    with tempfile.TemporaryDirectory(prefix=".lcms-", dir=destination.parent) as stage_dir:
        temporary = Path(stage_dir) / destination.name
        argv = [
            *[str(path) for path in sources],
            "--rt-tolerance", str(rt_tolerance),
            "--mz-tolerance", str(mz_tolerance),
            "--trend-threshold", str(trend_threshold),
            "--json-output", str(temporary),
            "--json-errors",
        ]
        if ignore_instrument:
            argv.append("--ignore-instrument")
        stderr = io.StringIO()
        captured_stdout = io.TextIOWrapper(io.BytesIO(), encoding="utf-8")
        try:
            try:
                with redirect_stdout(captured_stdout), redirect_stderr(stderr):
                    exit_code = main(argv)
            except SystemExit as exc:
                raise RuntimeError(
                    f"LCMS analysis exited unexpectedly with code {exc.code}"
                ) from exc
        finally:
            captured_stdout.close()
        if exit_code:
            raise RuntimeError(stderr.getvalue().strip() or "LCMS analysis failed")
        _assert_output(temporary)
        try:
            parsed = json.loads(temporary.read_text(encoding="utf-8"))
        except (json.JSONDecodeError, OSError) as exc:
            raise RuntimeError(f"LCMS analysis produced invalid JSON: {exc}") from exc
        if isinstance(parsed, dict):
            groups = [parsed]
        elif isinstance(parsed, list) and parsed and all(
            isinstance(group, dict) for group in parsed
        ):
            groups = parsed
        else:
            raise RuntimeError(
                "LCMS analysis JSON must be an object or a non-empty list of objects"
            )
        warnings = []
        for group in groups:
            warnings.extend(_normalize_warnings(group.get("warnings")))
        _commit_staged_outputs([(temporary, destination)])
    return _contract(
        {"json": str(destination)},
        warnings=warnings,
        metadata={"file_count": len(sources), "group_count": len(groups)},
    )


def assemble_lab_book(
    input_dir: str,
    experiment: Optional[str] = None,
    tracking_json: Optional[str] = None,
    output_path: Optional[str] = None,
) -> dict[str, Any]:
    """Assemble a deterministic lab-book entry from experiment files."""
    directory = _directory(input_dir)
    tracking = str(_source(tracking_json, suffixes=(".json",))) if tracking_json else None
    anchor = directory / "lab-book.txt"
    destination = _destination(anchor, output_path, tag="assembled", suffix=".txt")
    from cdxml_toolkit.analysis.deterministic.procedure_writer import main

    with tempfile.TemporaryDirectory(prefix=".lab-book-", dir=destination.parent) as stage_dir:
        temporary = Path(stage_dir) / destination.name
        argv = ["--input-dir", str(directory), "--output", str(temporary), "--json-errors"]
        if experiment:
            argv.extend(["--experiment", experiment])
        if tracking:
            argv.extend(["--tracking-json", tracking])
        with redirect_stdout(io.StringIO()), redirect_stderr(io.StringIO()) as stderr:
            try:
                exit_code = main(argv)
            except SystemExit as exc:
                raise RuntimeError(
                    f"Lab-book assembly exited unexpectedly with code {exc.code}"
                ) from exc
        if exit_code:
            raise RuntimeError(stderr.getvalue().strip() or "Lab-book assembly failed")
        _assert_output(temporary)
        _commit_staged_outputs([(temporary, destination)])
    return _contract({"text": str(destination)}, metadata={"experiment": experiment})


def parse_scifinder_rdf(
    input_path: str,
    resolve_cas: bool = False,
    output_path: Optional[str] = None,
    confirm_pubchem: bool = False,
) -> dict[str, Any]:
    """Parse SciFinder RDF and optionally enrich CAS data over the network."""
    if resolve_cas and not confirm_pubchem:
        raise ValueError(
            "resolve_cas uses PubChem network lookups; set confirm_pubchem=True "
            "to confirm this explicitly"
        )
    source = _source(input_path, suffixes=(".rdf",))
    from cdxml_toolkit.perception.rdf_parser import parse_rdf, reaction_to_dict, resolve_cas_numbers

    reactions = list(parse_rdf(str(source)))
    if not reactions:
        raise ValueError("SciFinder RDF contained no reactions")
    before_records = [
        copy.deepcopy(reaction_to_dict(reaction)) for reaction in reactions
    ]
    if resolve_cas:
        for reaction in reactions:
            resolve_cas_numbers(reaction)
    records = [reaction_to_dict(reaction) for reaction in reactions]
    cas_values = []

    def collect_cas(value: Any) -> None:
        if isinstance(value, dict):
            for key, child in value.items():
                if key == "cas" and isinstance(child, str) and child.strip():
                    cas_values.append(child.strip())
                collect_cas(child)
        elif isinstance(value, list):
            for child in value:
                collect_cas(child)

    collect_cas(records)
    def resolution_candidates(values: list[Any]) -> list[dict[str, Any]]:
        candidates = []
        for record in values:
            if not isinstance(record, dict):
                continue
            for variation in record.get("variations", []):
                if not isinstance(variation, dict):
                    continue
                for category in ("reagents", "catalysts", "solvents"):
                    for entry in variation.get(category, []):
                        if (
                            isinstance(entry, dict)
                            and isinstance(entry.get("cas"), str)
                            and entry["cas"].strip()
                        ):
                            candidates.append(entry)
        return candidates

    before_candidates = resolution_candidates(before_records)
    after_candidates = resolution_candidates(records)
    cas_resolutions = []
    if resolve_cas:
        for index, before in enumerate(before_candidates):
            after = after_candidates[index] if index < len(after_candidates) else {}
            fields_added = [
                key
                for key in ("name", "mw", "formula", "smiles")
                if before.get(key) in (None, "") and after.get(key) not in (None, "")
            ]
            resolution = {
                "cas": before["cas"].strip(),
                "status": "resolved" if fields_added else "unresolved",
                "fields_added": fields_added,
            }
            if fields_added and after.get("source"):
                resolution["source"] = after["source"]
            cas_resolutions.append(resolution)

    resolution_candidates_after = after_candidates
    attempted_cas_resolution_count = len(before_candidates) if resolve_cas else 0
    resolved_cas_count = sum(
        resolution["status"] == "resolved" for resolution in cas_resolutions
    )
    unresolved_cas_count = attempted_cas_resolution_count - resolved_cas_count
    destination = _destination(source, output_path, tag="parsed", suffix=".json")
    payload = _contract(
        {"json": str(destination)},
        metadata={
            "reaction_count": len(records),
            "resolve_cas": resolve_cas,
            "confirm_pubchem": confirm_pubchem,
            "cas_count": len(cas_values),
            "unique_cas_count": len(set(cas_values)),
            "cas_resolution_candidate_count": len(resolution_candidates_after),
            "attempted_cas_resolution_count": attempted_cas_resolution_count,
            "resolved_cas_count": resolved_cas_count,
            "unresolved_cas_count": unresolved_cas_count,
            "cas_resolutions": cas_resolutions,
        },
    )
    with tempfile.TemporaryDirectory(prefix=".rdf-", dir=destination.parent) as stage_dir:
        temporary = Path(stage_dir) / destination.name
        temporary.write_text(
            json.dumps({"reactions": records}, indent=2), encoding="utf-8"
        )
        _assert_output(temporary)
        _commit_staged_outputs([(temporary, destination)])
    return artifact_safety.with_artifacts(payload, [destination])


def segment_large_scheme(
    cdxml_path: str,
    output_path: Optional[str] = None,
    verbose: bool = False,
) -> dict[str, Any]:
    """Segment a disconnected or multi-panel CDXML scheme into logical regions."""
    source = _source(cdxml_path, suffixes=(".cdxml",))
    destination = _destination(source, output_path, tag="segments", suffix=".json")
    from cdxml_toolkit.perception.scheme_segmenter import classify_scheme_complexity, segment_scheme

    result = segment_scheme(str(source), verbose=verbose)
    segment_count = result.num_segments
    if not isinstance(segment_count, int) or isinstance(segment_count, bool) or segment_count <= 0:
        raise RuntimeError("Scheme segmentation produced no segments")
    data = result.to_dict()
    with tempfile.TemporaryDirectory(prefix=".segments-", dir=destination.parent) as stage_dir:
        temporary = Path(stage_dir) / destination.name
        temporary.write_text(json.dumps(data, indent=2), encoding="utf-8")
        _assert_output(temporary)
        _commit_staged_outputs([(temporary, destination)])
    return _contract(
        {"json": str(destination)},
        metadata={"complexity": classify_scheme_complexity(str(source)), "segments": segment_count},
    )


PUBLIC_TOOLS = {
    fn.__name__: fn
    for fn in (
        clean_scheme_layout, merge_reaction_schemes, polish_reaction_scheme,
        render_cdxml_files, fill_office_template, batch_embed_cdxml_in_office,
        inspect_chemdraw_objects_in_office, replace_chemdraw_objects_in_office,
        discover_experiment_files, analyze_lcms_series, assemble_lab_book,
        parse_scifinder_rdf, segment_large_scheme, diagnose_runtime,
    )
}
